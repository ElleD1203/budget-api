from fastapi import FastAPI, HTTPException, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.templating import Jinja2Templates
from pathlib import Path
from pydantic import BaseModel
from openpyxl import load_workbook
from datetime import datetime, date as date_type
from transaction_data import transactions 

app = FastAPI()
BASE_DIR = Path(__file__).resolve().parent
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class TransactionCreate(BaseModel):
    date: str
    merchant: str
    amount: float

class TransactionUpdate(BaseModel):
    date: str | None = None
    merchant: str | None = None
    amount: float | None = None


def get_transaction_type(amount: float):
    if amount > 0:
        return "income"
    elif amount < 0:
        return "expense"
    else:
        return "zero"

def get_next_id():
    if len(transactions) == 0:
        return 1

    max_id = max(transaction["id"] for transaction in transactions)
    return max_id + 1 

def clean_date(value):
    """
    convert excel date values into yyyy-mm-dd text.
    """
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    
    if isinstance(value, date_type):
        return value.strftime("%Y-%m-%d")
    
    return str(value).strip()

def clean_amount(value):
    """
    convert excel amount value into a number
    """

    try:
        return float(value)
    except:
        raise HTTPException(
            status_code=400,
            detail=f"invalid amount value:{value}"
        )

@app.get("/")
def home():
    return {"message": "Budget API is running"}


@app.get("/api/v1/transactions")
def get_transactions():
    return transactions

@app.get("/transactions-page")
def transactions_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="transactions.html",
        context={
            "transactions": transactions
        }
    )

@app.post("/api/v1/transactions")
def create_transaction(transaction: TransactionCreate):
    new_transaction = {
        "id": get_next_id(),
        "date": transaction.date,
        "merchant": transaction.merchant,
        "amount": transaction.amount,
        "transaction_type": get_transaction_type(transaction.amount),
    }

    transactions.append(new_transaction)

    return new_transaction

@app.patch("/api/v1/transactions/{transaction_id}")
def update_transaction(transaction_id: int, updated_transaction: TransactionUpdate):
    for transaction in transactions:
        if transaction["id"] == transaction_id:

            if updated_transaction.date is not None:
                transaction["date"] = updated_transaction.date

            if updated_transaction.merchant is not None:
                transaction["merchant"] = updated_transaction.merchant

            if updated_transaction.amount is not None:
                transaction["amount"] = updated_transaction.amount
                transaction["transaction_type"] = get_transaction_type(updated_transaction.amount)

            return transaction

    raise HTTPException(status_code=404, detail="Transaction not found")

@app.delete("/api/v1/transactions/{transaction_id}")
def delete_transaction(transaction_id: int):
    for transaction in transactions:
        if transaction["id"] == transaction_id:
            transactions.remove(transaction)

            return {
                "message": "Transaction deleted successfully",
                "deleted_transaction": transaction,
            }

    raise HTTPException(status_code=404, detail="Transaction not found")

@app.get("/api/v1/summary")
def get_summary(date: str | None = None):
    total_income = 0
    total_expense = 0

    for transaction in transactions:
        if date is not None and transaction["date"] != date:
            continue

        amount = transaction["amount"]

        if amount > 0:
            total_income = total_income + amount
        elif amount < 0:
            total_expense = total_expense + amount

    net_amount = total_income + total_expense

    return {
        "date": date if date is not None else "all",
        "total_income": total_income,
        "total_expense": total_expense,
        "net_amount": net_amount,
    }

@app.post("/api/v1/transactions/upload-excel")
def upload_excel(file:UploadFile = File(...)):
    """
    upload an excel file and add its rows to the transactions list.

    require excel columns:
    - date
    - merchant
    - amount
    """

    if not file.filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail= "only .xlsx excel file are supported"
        )

    workbook = load_workbook(file.file, data_only=True)
    sheet = workbook.active

    # read header row
    headers = []
    for cell in sheet[1]:
        if cell.value is not None:
            headers.append(str(cell.value).strip().lower())
        else:
            headers.append("")

    required_columns = ["date","merchant", "amount"]

    for column in required_columns:
        if column not in headers:
                raise HTTPException(
                    status_code=400,
                    detail=f"missing required column:{column}"
                )

    date_index = headers.index("date")
    merchant_index = headers.index("merchant")
    amount_index = headers.index("amount")

    uploaded_transactions = []

 # Read data rows, starting from row 2
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_date = row[date_index]
        row_merchant = row[merchant_index]
        row_amount = row[amount_index]

        # Skip completely empty rows
        if row_date is None and row_merchant is None and row_amount is None:
            continue

        if row_date is None or row_merchant is None or row_amount is None:
            raise HTTPException(
                status_code=400,
                detail="Each row must have date, merchant, and amount"
            )

        amount = clean_amount(row_amount)

        new_transaction = {
            "id": get_next_id(),
            "date": clean_date(row_date),
            "merchant": str(row_merchant).strip(),
            "amount": amount,
            "transaction_type": get_transaction_type(amount),
        }

        transactions.append(new_transaction)
        uploaded_transactions.append(new_transaction)

    return {
        "message": "Excel file uploaded successfully",
        "rows_added": len(uploaded_transactions),
        "transactions_added": uploaded_transactions,
    }

"""
add a comparison chart 
"""

@app.get("/api/v1/summary/chart")
def get_summary_chart(date: str | None = None):
    total_positive = 0
    total_negative = 0

    for transaction in transactions:
        if date is not None and transaction["date"] != date:
            continue

        amount = transaction["amount"]

        if amount > 0:
            total_positive = total_positive + amount
        elif amount < 0:
            total_negative = total_negative + amount

    return {
        "date": date if date is not None else "all",
        "total_positive": round(total_positive, 2),
        "total_negative": round(total_negative, 2),
        "total_negative_for_chart": round(abs(total_negative), 2),
    }
